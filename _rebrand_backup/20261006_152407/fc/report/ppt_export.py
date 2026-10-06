"""Ekspor Executive Report ke PowerPoint (BRD 2.5)."""
from datetime import datetime
from io import BytesIO
from typing import List, Optional, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ..config import Targets
from ..intel.alert_impact import compute_alerts
from ..intel.narrative import generate_narrative
from ..intel.period_compare import available_periods, compare_periods
from ..intel.recommend import generate_recommendations
from ..kpi import Scope, inventory_table, summarize
from ..pipeline import Dataset


# ==== WARNA TEMA ====
COLOR_PRIMARY = RGBColor(0x1E, 0x3A, 0x8A)     # Biru tua
COLOR_ACCENT = RGBColor(0x3B, 0x82, 0xF6)      # Biru muda
COLOR_DARK = RGBColor(0x0F, 0x17, 0x2A)        # Navy
COLOR_LIGHT = RGBColor(0xF8, 0xFA, 0xFC)       # Abu terang
COLOR_RED = RGBColor(0xDC, 0x26, 0x26)
COLOR_YELLOW = RGBColor(0xEA, 0xB3, 0x08)
COLOR_GREEN = RGBColor(0x16, 0xA3, 0x4A)
COLOR_GRAY = RGBColor(0x64, 0x74, 0x8B)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _add_title_slide(prs: Presentation, ds: Dataset, scope: Scope) -> None:
    """Slide 1 — Cover."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Background biru tua
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_PRIMARY
    bg.line.fill.background()

    # Judul
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(8.4), Inches(1.5))
    tf = tb.text_frame
    tf.text = "EXECUTIVE REPORT"
    p = tf.paragraphs[0]
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(8.4), Inches(0.8))
    tf2 = tb2.text_frame
    tf2.text = "AI Factory Controller"
    p2 = tf2.paragraphs[0]
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(0x60, 0xA5, 0xFA)
    p2.alignment = PP_ALIGN.CENTER

    # Info
    tb3 = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(8.4), Inches(1.5))
    tf3 = tb3.text_frame
    lines = [
        f"Sumber: {ds.report.source_label}",
        f"Dibaca: {ds.report.loaded_at}",
        f"Periode: {scope.start or '-'} s/d {scope.end or '-'}",
        f"Plant: {scope.plant or 'Semua'} | Line: {scope.line or 'Semua'}",
    ]
    for i, line in enumerate(lines):
        if i == 0:
            tf3.text = line
            p = tf3.paragraphs[0]
        else:
            p = tf3.add_paragraph()
            p.text = line
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
        p.alignment = PP_ALIGN.CENTER

    # Banner DEMO
    if ds.is_demo:
        tb4 = slide.shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(8.4), Inches(0.5))
        tf4 = tb4.text_frame
        tf4.text = "*** DEMO DATA - bukan data pabrik sebenarnya ***"
        p4 = tf4.paragraphs[0]
        p4.font.size = Pt(14)
        p4.font.bold = True
        p4.font.color.rgb = RGBColor(0xFB, 0xBF, 0x24)
        p4.alignment = PP_ALIGN.CENTER


def _add_kpi_slide(prs: Presentation, s, targets: Targets) -> None:
    """Slide 2 — Ringkasan KPI."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _add_slide_header(slide, "Ringkasan KPI")

    def _fmt(kpi, fmt):
        return fmt.format(kpi.value) if kpi.available else "—"

    kpis = [
        ("Output", _fmt(s.output_kg, "{:,.0f} Kg")),
        ("Yield", _fmt(s.yield_pct, "{:.2f}%")),
        ("Scrap", _fmt(s.scrap_pct, "{:.2f}%")),
        ("OEE", _fmt(s.oee.oee, "{:.2f}%")),
        ("COGM", _fmt(s.cogm, "Rp {:,.0f}")),
        ("Cost/Kg", _fmt(s.cost_per_kg, "Rp {:,.0f}")),
        ("Score",
         f"{s.score.score}/100" if s.score.score is not None else "—"),
        ("Kategori", s.score.category),
    ]

    # Grid 4x2
    x_start = Inches(0.6)
    y_start = Inches(1.8)
    card_w = Inches(2.15)
    card_h = Inches(1.5)
    gap_x = Inches(0.15)
    gap_y = Inches(0.2)

    for i, (label, value) in enumerate(kpis):
        row = i // 4
        col = i % 4
        x = x_start + (card_w + gap_x) * col
        y = y_start + (card_h + gap_y) * row

        # Card background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_LIGHT
        card.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)

        # Label
        tb = slide.shapes.add_textbox(x + Inches(0.1), y + Inches(0.1), card_w, Inches(0.4))
        tf = tb.text_frame
        tf.text = label.upper()
        tf.paragraphs[0].font.size = Pt(9)
        tf.paragraphs[0].font.color.rgb = COLOR_GRAY
        tf.paragraphs[0].font.bold = True

        # Value
        tb2 = slide.shapes.add_textbox(x + Inches(0.1), y + Inches(0.55), card_w, Inches(0.9))
        tf2 = tb2.text_frame
        tf2.text = value
        tf2.paragraphs[0].font.size = Pt(20)
        tf2.paragraphs[0].font.bold = True
        tf2.paragraphs[0].font.color.rgb = COLOR_DARK

    # Target info
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(6.0), Inches(9.0), Inches(0.5))
    tf = tb.text_frame
    tf.text = (f"Target aktif: Yield ≥ {targets.yield_min:g}% · "
               f"Scrap ≤ {targets.scrap_max:g}% · OEE ≥ {targets.oee_min:g}%")
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = COLOR_GRAY
    tf.paragraphs[0].font.italic = True


def _add_alert_slide(prs: Presentation, alerts: List, s) -> None:
    """Slide 3 — Peringatan Aktif dengan impact Rp."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_slide_header(slide, "🚨 Peringatan Aktif (berdasarkan dampak Rp)")

    if not alerts:
        tb = slide.shapes.add_textbox(Inches(0.6), Inches(3.0), Inches(9.0), Inches(1.0))
        tf = tb.text_frame
        tf.text = "✅ Tidak ada peringatan aktif. Semua KPI dalam batas target."
        tf.paragraphs[0].font.size = Pt(20)
        tf.paragraphs[0].font.color.rgb = COLOR_GREEN
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        return

    # Total impact
    total_annual = sum(abs(a.annual_rp) for a in alerts)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(9.0), Inches(0.6))
    tf = tb.text_frame
    tf.text = f"Total potensi dampak: Rp {total_annual:,.0f}/tahun"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = COLOR_RED

    # List alerts (max 3)
    y = Inches(2.3)
    for a in alerts[:3]:
        # Card
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.6), y, Inches(9.0), Inches(1.4)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_LIGHT
        card.line.color.rgb = COLOR_RED if a.severity == "high" else COLOR_YELLOW
        card.line.width = Pt(2)

        # Title
        tb = slide.shapes.add_textbox(Inches(0.8), y + Inches(0.1), Inches(8.6), Inches(0.4))
        tf = tb.text_frame
        icon = {"high": "🔴", "medium": "🟡", "info": "🔵"}.get(a.severity, "⚪")
        tf.text = f"{icon} {a.title}"
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = COLOR_DARK

        # Detail
        tb2 = slide.shapes.add_textbox(Inches(0.8), y + Inches(0.5), Inches(8.6), Inches(0.85))
        tf2 = tb2.text_frame
        tf2.text = f"{a.message} · Dampak: Rp {a.monthly_rp:,.0f}/bulan (Rp {a.annual_rp:,.0f}/tahun)"
        tf2.paragraphs[0].font.size = Pt(11)
        tf2.paragraphs[0].font.color.rgb = COLOR_GRAY

        p2 = tf2.add_paragraph()
        p2.text = f"Saran: {a.fix_hint}"
        p2.font.size = Pt(10)
        p2.font.italic = True
        p2.font.color.rgb = COLOR_GRAY

        y += Inches(1.55)


def _add_recommendation_slide(prs: Presentation, recs: List) -> None:
    """Slide 4 — Rekomendasi prioritas."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_slide_header(slide, "💡 Rekomendasi Aksi Prioritas")

    if not recs:
        tb = slide.shapes.add_textbox(Inches(0.6), Inches(3.0), Inches(9.0), Inches(1.0))
        tf = tb.text_frame
        tf.text = "Tidak ada rekomendasi saat ini."
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.color.rgb = COLOR_GRAY
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        return

    y = Inches(1.6)
    for i, r in enumerate(recs[:3], 1):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.6), y, Inches(9.0), Inches(1.6)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_LIGHT
        card.line.color.rgb = COLOR_ACCENT

        # Number badge
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.75), y + Inches(0.15), Inches(0.5), Inches(0.5)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_PRIMARY
        badge.line.fill.background()
        btf = badge.text_frame
        btf.text = str(i)
        btf.paragraphs[0].font.size = Pt(16)
        btf.paragraphs[0].font.bold = True
        btf.paragraphs[0].font.color.rgb = COLOR_WHITE
        btf.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Title
        tb = slide.shapes.add_textbox(Inches(1.4), y + Inches(0.1), Inches(6.5), Inches(0.4))
        tf = tb.text_frame
        tf.text = r.title
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = COLOR_DARK

        # Detail
        tb2 = slide.shapes.add_textbox(Inches(1.4), y + Inches(0.55), Inches(6.5), Inches(0.95))
        tf2 = tb2.text_frame
        tf2.text = r.context[:150]
        tf2.paragraphs[0].font.size = Pt(10)
        tf2.paragraphs[0].font.color.rgb = COLOR_GRAY

        # Impact
        impact = r.annual_rp if r.annual_rp > 0 else r.monthly_rp
        suffix = "/tahun" if r.annual_rp > 0 else " (1x)"
        tb3 = slide.shapes.add_textbox(Inches(7.9), y + Inches(0.3), Inches(1.6), Inches(1.0))
        tf3 = tb3.text_frame
        tf3.text = "HEMAT"
        tf3.paragraphs[0].font.size = Pt(9)
        tf3.paragraphs[0].font.color.rgb = COLOR_GRAY
        tf3.paragraphs[0].alignment = PP_ALIGN.RIGHT

        p2 = tf3.add_paragraph()
        if impact >= 1_000_000_000:
            val = f"Rp {impact/1_000_000_000:.2f} M"
        elif impact >= 1_000_000:
            val = f"Rp {impact/1_000_000:.1f} jt"
        else:
            val = f"Rp {impact:,.0f}"
        p2.text = val
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_GREEN
        p2.alignment = PP_ALIGN.RIGHT

        p3 = tf3.add_paragraph()
        p3.text = suffix
        p3.font.size = Pt(9)
        p3.font.color.rgb = COLOR_GRAY
        p3.alignment = PP_ALIGN.RIGHT

        y += Inches(1.7)


def _add_compare_slide(prs: Presentation, ds: Dataset) -> None:
    """Slide 5 — Perbandingan periode."""
    periods = available_periods(ds)
    if len(periods) < 2:
        return

    if len(periods) >= 3:
        prev, curr = periods[-3], periods[-2]
    else:
        prev, curr = periods[-2], periods[-1]

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_slide_header(slide, f"📅 Perbandingan Periode: {prev} vs {curr}")

    comps = compare_periods(ds, prev, curr)

    # Table header
    y = Inches(1.6)
    row_h = Inches(0.5)
    headers = ["KPI", prev, curr, "Delta", "Status"]
    cols_x = [Inches(0.6), Inches(3.0), Inches(4.8), Inches(6.6), Inches(8.0)]
    cols_w = [Inches(2.4), Inches(1.8), Inches(1.8), Inches(1.4), Inches(1.6)]

    for i, (h, x, w) in enumerate(zip(headers, cols_x, cols_w)):
        cell = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, x, y, w, row_h
        )
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        cell.line.fill.background()
        ctf = cell.text_frame
        ctf.text = h
        ctf.paragraphs[0].font.size = Pt(11)
        ctf.paragraphs[0].font.bold = True
        ctf.paragraphs[0].font.color.rgb = COLOR_WHITE
        ctf.paragraphs[0].alignment = PP_ALIGN.CENTER

    y += row_h

    for c in comps:
        improving = c.is_improving
        if improving is True:
            status_icon, status_color = "✅ Membaik", COLOR_GREEN
        elif improving is False:
            status_icon, status_color = "🔴 Memburuk", COLOR_RED
        else:
            status_icon, status_color = "➖ Stabil", COLOR_GRAY

        row_data = [
            c.label,
            c.fmt_value(c.prev_value),
            c.fmt_value(c.curr_value),
            c.fmt_delta(),
            status_icon,
        ]

        for i, (data, x, w) in enumerate(zip(row_data, cols_x, cols_w)):
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, row_h)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_LIGHT if i != 4 else COLOR_WHITE
            cell.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)

            ctf = cell.text_frame
            ctf.text = str(data)
            ctf.paragraphs[0].font.size = Pt(10)
            ctf.paragraphs[0].font.color.rgb = status_color if i == 4 else COLOR_DARK
            if i == 4:
                ctf.paragraphs[0].font.bold = True
            ctf.paragraphs[0].alignment = PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER

        y += row_h


def _add_narrative_slide(prs: Presentation, narrative) -> None:
    """Slide 6 — Narasi eksekutif (ringkas)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_slide_header(slide, "📰 Ringkasan Eksekutif")

    # Headline
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(9.0), Inches(0.6))
    tf = tb.text_frame
    tf.text = narrative.headline
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = COLOR_PRIMARY

    # Executive summary
    tb2 = slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(9.0), Inches(2.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.text = narrative.executive_summary.replace("**", "")
    tf2.paragraphs[0].font.size = Pt(12)
    tf2.paragraphs[0].font.color.rgb = COLOR_DARK

    # Key points (2 kolom)
    y = Inches(4.3)

    # Pencapaian
    tb3 = slide.shapes.add_textbox(Inches(0.6), y, Inches(4.5), Inches(2.5))
    tf3 = tb3.text_frame
    tf3.text = "🟢 PENCAPAIAN"
    tf3.paragraphs[0].font.size = Pt(12)
    tf3.paragraphs[0].font.bold = True
    tf3.paragraphs[0].font.color.rgb = COLOR_GREEN
    for item in narrative.achievements.items[:3]:
        p = tf3.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_DARK

    # Kritis
    tb4 = slide.shapes.add_textbox(Inches(5.2), y, Inches(4.5), Inches(2.5))
    tf4 = tb4.text_frame
    tf4.text = "🔴 KRITIS"
    tf4.paragraphs[0].font.size = Pt(12)
    tf4.paragraphs[0].font.bold = True
    tf4.paragraphs[0].font.color.rgb = COLOR_RED
    for item in narrative.criticals.items[:3]:
        p = tf4.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_DARK


def _add_slide_header(slide, title: str) -> None:
    """Header konsisten untuk semua slide."""
    # Garis atas
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(1.1)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    # Judul
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(9.0), Inches(0.7))
    tf = tb.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = COLOR_WHITE

    # Footer
    tb2 = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(9.0), Inches(0.4))
    tf2 = tb2.text_frame
    tf2.text = f"AI Factory Controller · {datetime.now():%d-%m-%Y %H:%M}"
    tf2.paragraphs[0].font.size = Pt(8)
    tf2.paragraphs[0].font.color.rgb = COLOR_GRAY
    tf2.paragraphs[0].alignment = PP_ALIGN.CENTER


def generate_ppt(ds: Dataset, scope: Scope,
                 targets: Optional[Targets] = None) -> bytes:
    """Generate PPT Executive Report. Return bytes."""
    from ..config import TARGETS as DEFAULT_TARGETS
    if targets is None:
        targets = DEFAULT_TARGETS

    s = summarize(ds, scope)
    alerts = compute_alerts(s, targets)
    narrative = generate_narrative(ds, s, targets)
    try:
        recs = generate_recommendations(ds, s, targets)
    except Exception:
        recs = []

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Build slides
    _add_title_slide(prs, ds, scope)
    _add_kpi_slide(prs, s, targets)
    _add_alert_slide(prs, alerts, s)
    _add_recommendation_slide(prs, recs)
    _add_compare_slide(prs, ds)
    _add_narrative_slide(prs, narrative)

    # Save to bytes
    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()